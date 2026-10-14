"""
Make Excel versions of the 200-row pilot labeling sheets, with dropdown menus.

Usage:  python make_pilot_xlsx.py
Output: annotation/pilot_annotator_A.xlsx, annotation/pilot_annotator_B.xlsx
agreement.py reads these directly (no need to convert back to CSV).
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ASPECTS = ["배송_delivery", "가격_price", "품질_quality", "포장_packaging", "서비스_service", "기타_other"]
FONT = "Arial"


def build(sheet_csv, out_path, who):
    df = pd.read_csv(sheet_csv, dtype=str, keep_default_na=False)
    df = df[df["pilot"].str.lower() == "true"].head(200)

    wb = Workbook()
    ws = wb.active
    ws.title = "Label here"
    headers = ["item", "text", "text_sentiment"] + ASPECTS + ["unclear", "note"]
    ws.append(headers)
    for r in df.itertuples():
        ws.append([int(r.item), r.text] + [""] * (len(headers) - 2))

    head_fill = PatternFill("solid", start_color="2E4A52")
    input_fill = PatternFill("solid", start_color="FFF2CC")  # yellow = fill in
    thin = Side(style="thin", color="BFBFBF")
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n = len(df) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font = Font(name=FONT, size=11)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=(c.column == 2 or c.column == len(headers)))
            if c.column >= 3:
                c.fill = input_fill
                if c.column < len(headers):
                    c.alignment = Alignment(horizontal="center", vertical="top")
    widths = [6, 70, 15] + [13] * len(ASPECTS) + [9, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"

    dv_s = DataValidation(type="list", formula1='"pos,neg,mixed,neutral"', allow_blank=True,
                          error="Choose pos, neg, mixed or neutral", errorTitle="Invalid value", showErrorMessage=True)
    dv_a = DataValidation(type="list", formula1='"pos,neg"', allow_blank=True,
                          error="Choose pos or neg, or leave empty", errorTitle="Invalid value", showErrorMessage=True)
    dv_u = DataValidation(type="list", formula1='"1"', allow_blank=True, showErrorMessage=True)
    for dv in (dv_s, dv_a, dv_u):
        ws.add_data_validation(dv)
    dv_s.add(f"C2:C{n}")
    dv_a.add(f"D2:{ws.cell(1, 3 + len(ASPECTS)).column_letter}{n}")
    u = ws.cell(1, 4 + len(ASPECTS)).column_letter
    dv_u.add(f"{u}2:{u}{n}")

    # Guide sheet: legend, progress counter, example row
    g = wb.create_sheet("How to label")
    lines = [
        (f"Pilot labeling sheet — annotator {who}", True),
        ("Fill only the YELLOW cells on the 'Label here' sheet. Use the dropdown in each cell.", False),
        ("Do not discuss your labels with the other annotator until both of you are finished.", False),
        ("", False),
        ("text_sentiment: pos = satisfied overall · neg = dissatisfied overall · mixed = clear praise AND complaint · neutral = no opinion", False),
        ("Aspect columns: pos or neg if the review talks about it; leave EMPTY if not mentioned", False),
        ("  delivery = speed, arrival, missing items · price = price, value for money · quality = the product itself", False),
        ("  packaging = box, wrapping, damage in transit · service = seller response, exchange, refund · other = anything else", False),
        ("unclear: choose 1 if the text has no meaning (e.g. random letters)", False),
        ("Full rules and 10 examples: annotation/LABELING_GUIDE.md", False),
        ("", False),
        ("Example (not part of the data):", True),
    ]
    for i, (t, b) in enumerate(lines, 1):
        g.cell(i, 1, t).font = Font(name=FONT, bold=b, size=13 if i == 1 else 11)
    ex_head = ["text", "text_sentiment", "배송_delivery", "품질_quality"]
    ex_row = ["배송은 빠른데 사이즈가 너무 작아요", "mixed", "pos", "neg"]
    for j, (h, v) in enumerate(zip(ex_head, ex_row), 1):
        g.cell(13, j, h).font = Font(name=FONT, bold=True)
        g.cell(14, j, v).font = Font(name=FONT)
    g.cell(16, 1, "Rows labeled so far (out of 200):").font = Font(name=FONT, bold=True)
    g.cell(16, 2, f"=COUNTA('Label here'!C2:C{n})").font = Font(name=FONT, bold=True)
    g.column_dimensions["A"].width = 45
    for col in "BCD":
        g.column_dimensions[col].width = 16
    wb.save(out_path)
    print(f"Saved {out_path} ({len(df)} rows)")


if __name__ == "__main__":
    d = sys.argv[1] if len(sys.argv) > 1 else "annotation"
    build(f"{d}/annotator_A.csv", f"{d}/pilot_annotator_A.xlsx", "A")
    build(f"{d}/annotator_B.csv", f"{d}/pilot_annotator_B.xlsx", "B")
